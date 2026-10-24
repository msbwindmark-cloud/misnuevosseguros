from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse_lazy
from django.views.generic import ListView, CreateView, UpdateView, DeleteView
from django.contrib.admin.views.decorators import staff_member_required
from django.http import HttpResponse, HttpResponseNotFound, JsonResponse, FileResponse
from django.utils.decorators import method_decorator
from django.views.decorators.csrf import csrf_exempt
from django.contrib import messages
from django.views import View
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.units import cm
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl import Workbook
from .models import Banco, Benificiario, EnvioBanco
from .forms import BancoForm, BenificiarioForm, EnvioBancoForm


class BancoListView(ListView):
    model = Banco
    template_name = 'banco/lista.html'
    context_object_name = 'bancos'


class BenificiarioListView(ListView):
    model = Benificiario
    template_name = 'banco/beneficiarios.html'
    context_object_name = 'beneficiarios'


class EnvioBancoListView(ListView):
    model = EnvioBanco
    template_name = 'banco/envios.html'
    context_object_name = 'envios'


@method_decorator(staff_member_required, name='dispatch')
class BancoCreateView(CreateView):
    model = Banco
    form_class = BancoForm
    template_name = 'banco/banco_form.html'
    success_url = reverse_lazy('banco_list')


@method_decorator(staff_member_required, name='dispatch')
class BancoUpdateView(UpdateView):
    model = Banco
    form_class = BancoForm
    template_name = 'banco/banco_form.html'
    success_url = reverse_lazy('banco_list')


@method_decorator(staff_member_required, name='dispatch')
class BancoDeleteView(DeleteView):
    model = Banco
    template_name = 'banco/banco_confirm_delete.html'
    success_url = reverse_lazy('banco_list')


@method_decorator(staff_member_required, name='dispatch')
class BenificiarioCreateView(CreateView):
    model = Benificiario
    form_class = BenificiarioForm
    template_name = 'banco/beneficiario_form.html'
    success_url = reverse_lazy('benificiario_list')


@method_decorator(staff_member_required, name='dispatch')
class BenificiarioUpdateView(UpdateView):
    model = Benificiario
    form_class = BenificiarioForm
    template_name = 'banco/beneficiario_form.html'
    success_url = reverse_lazy('benificiario_list')


@method_decorator(staff_member_required, name='dispatch')
class BenificiarioDeleteView(DeleteView):
    model = Benificiario
    template_name = 'banco/beneficiario_confirm_delete.html'
    success_url = reverse_lazy('benificiario_list')


@method_decorator(staff_member_required, name='dispatch')
class EnvioBancoCreateView(View):
    template_name = 'banco/envio_form.html'
    success_url = reverse_lazy('envio_banco_list')

    def get(self, request):
        form = EnvioBancoForm()
        bancos = Banco.objects.all()
        beneficiarios = Benificiario.objects.all()
        return render(request, self.template_name, {
            'form': form, 'bancos': bancos, 'beneficiarios': beneficiarios
        })

    def post(self, request):
        data = request.POST
        files = request.FILES
        benificiario_id = data.get('benificiario')
        banco_id = data.get('banco')
        try:
            ben = Benificiario.objects.get(pk=int(benificiario_id)) if benificiario_id else None
        except (Benificiario.DoesNotExist, ValueError, TypeError):
            ben = None
        try:
            ban = Banco.objects.get(pk=int(banco_id)) if banco_id else None
        except (Banco.DoesNotExist, ValueError, TypeError):
            ban = None

        envio = EnvioBanco.objects.create(
            nom_envoie=ben.nom_benificiaire if ben else '',
            banque_envoie=ben.banque_benificiaire if ben else '',
            compte_envoie=ben.compte_benificiaire if ben else '',
            date_envoie=data.get('date_envoie') or None,
            montant_envoie=data.get('montant_envoie', 0),
            motif_envoie=data.get('motif_envoie', ''),
            benificiario=ben,
            banco=ban,
        )

        file = files.get('fileupload')
        if file:
            ext = file.name.split('.')[-1].lower() if '.' in file.name else ''
            import os
            from django.conf import settings
            upload_dir = settings.MEDIA_ROOT / 'uploads_banco'
            upload_dir.mkdir(parents=True, exist_ok=True)
            filepath = upload_dir / file.name
            with open(str(filepath), 'wb+') as dest:
                for chunk in file.chunks():
                    dest.write(chunk)
            envio.tipo_documento_envoie_banque = ext
            envio.nombre_documento_envoie_banque = file.name
            envio.tamanio_documento_envoie_banque = file.size
            envio.save()

        messages.success(request, f'Virement enregistré correctement')
        return redirect(self.success_url)


@method_decorator(staff_member_required, name='dispatch')
class EnvioBancoUpdateView(UpdateView):
    model = EnvioBanco
    form_class = EnvioBancoForm
    template_name = 'banco/envio_form.html'
    success_url = reverse_lazy('envio_banco_list')

    def form_valid(self, form):
        response = super().form_valid(form)
        envio = self.object
        file = self.request.FILES.get('fileupload')
        if file:
            ext = file.name.split('.')[-1].lower() if '.' in file.name else ''
            import os
            from django.conf import settings
            upload_dir = settings.MEDIA_ROOT / 'uploads_banco'
            upload_dir.mkdir(parents=True, exist_ok=True)
            filepath = upload_dir / file.name
            with open(str(filepath), 'wb+') as dest:
                for chunk in file.chunks():
                    dest.write(chunk)
            envio.tipo_documento_envoie_banque = ext
            envio.nombre_documento_envoie_banque = file.name
            envio.tamanio_documento_envoie_banque = file.size
            envio.save()
        return response


@method_decorator(staff_member_required, name='dispatch')
class EnvioBancoDeleteView(DeleteView):
    model = EnvioBanco
    template_name = 'banco/envio_confirm_delete.html'
    success_url = reverse_lazy('envio_banco_list')


@staff_member_required
def envio_banco_last_pdf(request):
    envio = EnvioBanco.objects.first()
    if not envio:
        return HttpResponse('No hay virements', status=404)
    return envio_banco_pdf(request, envio.pk)


@staff_member_required
def envio_banco_pdf(request, pk):
    from num2words import num2words
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas as canvasmod
    from reportlab.lib.units import mm
    from reportlab.lib import colors as rl_colors

    try:
        envio = EnvioBanco.objects.get(pk=pk)
    except EnvioBanco.DoesNotExist:
        return HttpResponse('Envío no encontrado', status=404)

    banco = envio.banco
    fecha = envio.date_envoie.strftime('%d/%m/%Y') if envio.date_envoie else ''
    montant = float(envio.montant_envoie)
    parte_entera = int(montant)
    parte_decimal = int(round((montant - parte_entera) * 100))
    monto_palabras = num2words(parte_entera, lang='fr').upper()
    monto_centimos = num2words(parte_decimal, lang='fr').upper() if parte_decimal > 0 else 'ZERO'

    buffer = BytesIO()
    w, h = A4
    c = canvasmod.Canvas(buffer, pagesize=A4)
    c.setTitle('Envoie banque')
    c.setAuthor('SIS')

    verde = rl_colors.Color(2/255, 157/255, 116/255)
    texto_verde = rl_colors.Color(240/255, 255/255, 240/255)
    x = 15 * mm

    c.setFont('Helvetica', 12)
    c.drawString(x, h - 60 * mm, f'Rabat, le  : {fecha}')

    c.setFont('Helvetica', 12)
    c.drawString(x, h - 70 * mm, str(banco.nom_banque))
    c.drawString(x, h - 78 * mm, str(banco.distrit_banque))
    c.drawString(x, h - 86 * mm, str(banco.ville_banque))

    c.setFont('Helvetica', 12)
    c.drawString(x, h - 115 * mm, 'Messieurs,')

    c.setFont('Helvetica', 11)
    intro = (
        f'Par le debit de notre compte n {banco.compte_banque}, '
        f'nous vous prions de bien vouloir virer les montants aux comptes '
        f'des beneficiaires dont la liste es indiquee ci-apres:'
    )
    c.drawString(x, h - 123 * mm, intro)

    y_table = h - 145 * mm
    w_nom = 50 * mm
    w_banque = 40 * mm
    w_compte = 60 * mm
    w_montant = 30 * mm
    row_h = 8 * mm

    c.setFillColor(verde)
    c.rect(x, y_table - row_h, w_nom, row_h, fill=1, stroke=1)
    c.rect(x + w_nom, y_table - row_h, w_banque, row_h, fill=1, stroke=1)
    c.rect(x + w_nom + w_banque, y_table - row_h, w_compte, row_h, fill=1, stroke=1)
    c.rect(x + w_nom + w_banque + w_compte, y_table - row_h, w_montant, row_h, fill=1, stroke=1)

    c.setFillColor(texto_verde)
    c.setFont('Helvetica-Bold', 12)
    c.drawString(x + 2 * mm, y_table - row_h + 2 * mm, 'Noms')
    c.drawString(x + w_nom + 2 * mm, y_table - row_h + 2 * mm, 'Banque')
    c.drawString(x + w_nom + w_banque + 2 * mm, y_table - row_h + 2 * mm, 'N Compte')
    c.drawString(x + w_nom + w_banque + w_compte + 2 * mm, y_table - row_h + 2 * mm, 'Montant')

    y_data = y_table - row_h * 2
    c.setFillColor(rl_colors.black)
    c.setFont('Helvetica', 11)
    c.rect(x, y_data, w_nom, row_h, fill=0, stroke=1)
    c.rect(x + w_nom, y_data, w_banque, row_h, fill=0, stroke=1)
    c.rect(x + w_nom + w_banque, y_data, w_compte, row_h, fill=0, stroke=1)
    c.rect(x + w_nom + w_banque + w_compte, y_data, w_montant, row_h, fill=0, stroke=1)

    c.drawString(x + 2 * mm, y_data + 2 * mm, str(envio.nom_envoie))
    c.drawString(x + w_nom + 2 * mm, y_data + 2 * mm, str(envio.banque_envoie))
    c.drawCentredString(x + w_nom + w_banque + w_compte / 2, y_data + 2 * mm, str(envio.compte_envoie))
    c.drawCentredString(x + w_nom + w_banque + w_compte + w_montant / 2, y_data + 2 * mm, f'{montant:.2f}')

    y_total = y_data - row_h
    c.setFont('Helvetica-Bold', 11)
    c.rect(x, y_total, w_nom + w_banque + w_compte, row_h, fill=0, stroke=1)
    c.rect(x + w_nom + w_banque + w_compte, y_total, w_montant, row_h, fill=0, stroke=1)
    total_label_w = w_nom + w_banque + w_compte
    c.drawRightString(x + total_label_w - 2 * mm, y_total + 2 * mm, 'Total General')
    c.drawCentredString(x + w_nom + w_banque + w_compte + w_montant / 2, y_total + 2 * mm, f'{montant:.2f}')

    y_after_table = y_total - 12 * mm
    c.setFont('Helvetica-Bold', 12)
    c.drawString(x, y_after_table, f'Motif Virement: {envio.motif_envoie}.')

    y_montant_text = y_after_table - 10 * mm
    texto_virement = (
        f'ARRETE LE PRESENT ORDRE DE VIREMENT A LA SOMME DE '
        f'{monto_palabras} DHS ET {monto_centimos} CTS.'
    )
    c.setFont('Helvetica-Bold', 12)
    text_obj = c.beginText(x, y_montant_text)
    text_obj.setLeading(14)
    words = texto_virement.split(' ')
    line = ''
    for word in words:
        test = line + ' ' + word if line else word
        if c.stringWidth(test, 'Helvetica-Bold', 12) > (w - 2 * x):
            text_obj.textLine(line)
            line = word
        else:
            line = test
    if line:
        text_obj.textLine(line)
    c.drawText(text_obj)

    y_salut = y_montant_text - 20 * mm
    c.setFont('Helvetica', 12)
    c.drawString(x, y_salut, 'Veuillez agreer, Messieurs, nos salutations distinguees.')

    from datetime import datetime
    c.setFont('Helvetica', 9)
    c.drawString(x, y_salut - 10 * mm, f'Document: Ordre de virement N {envio.id}')
    c.drawString(x, y_salut - 14 * mm, f'Imprime le: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}')

    c.save()
    buffer.seek(0)
    nombre = f"ordre_virement_{envio.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return HttpResponse(buffer, content_type='application/pdf', headers={
        'Content-Disposition': f'attachment; filename="{nombre}"'
    })


@staff_member_required
def envio_banco_list_pdf(request):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    elements = []
    elements.append(Paragraph('Lista de Envíos Bancarios', styles['Title']))
    elements.append(Spacer(1, 0.5*cm))
    data = [['ID', 'Nombre', 'Banco', 'Cuenta', 'Fecha', 'Monto', 'Beneficiario']]
    for e in EnvioBanco.objects.all():
        fecha = e.date_envoie.strftime('%d/%m/%Y') if e.date_envoie else ''
        data.append([str(e.id), e.nom_envoie, e.banque_envoie, e.compte_envoie, fecha, str(e.montant_envoie), str(e.benificiario)])
    table = Table(data, colWidths=[1.5*cm, 3*cm, 3*cm, 3*cm, 2.5*cm, 2.5*cm, 3.5*cm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4e73df')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fc')]),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    return HttpResponse(buffer, content_type='application/pdf', headers={
        'Content-Disposition': 'attachment; filename="envios_bancarios.pdf"'
    })


@staff_member_required
def envio_banco_download(request, pk):
    from django.conf import settings
    import os
    envio = get_object_or_404(EnvioBanco, pk=pk)
    fname = getattr(envio, 'nombre_documento_envoie_banque', '') or ''
    if not fname:
        return HttpResponseNotFound('No hay documento asociado con este virement.')
    
    path = os.path.join(str(settings.MEDIA_ROOT), 'uploads_banco', fname)
    
    if not os.path.isfile(path):
        return HttpResponseNotFound('El fichero no existe en el servidor.')
    
    # Abrimos el fichero sin el 'with' para que FileResponse lo lea y cierre correctamente al terminar
    f = open(path, 'rb')
    return FileResponse(f, as_attachment=True, filename=os.path.basename(fname))


@staff_member_required
def export_envios_banco_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Envíos Bancarios'
    headers = ['ID', 'Nombre', 'Banco', 'Cuenta', 'Fecha', 'Monto', 'Motivo', 'Beneficiario', 'Banco (FK)', 'Creado']
    header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
    header_font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal='center')
    for row, e in enumerate(EnvioBanco.objects.all().order_by('-id'), 2):
        ws.cell(row=row, column=1, value=e.id)
        ws.cell(row=row, column=2, value=e.nom_envoie)
        ws.cell(row=row, column=3, value=e.banque_envoie)
        ws.cell(row=row, column=4, value=e.compte_envoie)
        ws.cell(row=row, column=5, value=e.date_envoie.strftime('%d/%m/%Y') if e.date_envoie else '')
        ws.cell(row=row, column=6, value=float(e.montant_envoie))
        ws.cell(row=row, column=7, value=e.motif_envoie)
        ws.cell(row=row, column=8, value=str(e.benificiario))
        ws.cell(row=row, column=9, value=str(e.banco))
        ws.cell(row=row, column=10, value=e.creado_el.strftime('%d/%m/%Y %H:%M') if e.creado_el else '')
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 22
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 30
    ws.column_dimensions['H'].width = 25
    ws.column_dimensions['I'].width = 20
    ws.column_dimensions['J'].width = 18
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="envios_bancarios.xlsx"'
    wb.save(response)
    return response


@staff_member_required
def get_beneficiario_json(request):
    pk = request.GET.get('pk', '')
    if not pk or not pk.isdigit():
        return JsonResponse({'ok': False})
    try:
        b = Benificiario.objects.get(pk=int(pk))
        return JsonResponse({'ok': True, 'nom': b.nom_benificiaire, 'banque': b.banque_benificiaire, 'compte': b.compte_benificiaire})
    except Benificiario.DoesNotExist:
        return JsonResponse({'ok': False})


# --- AJAX Banco (otro_banque.php) ---

@csrf_exempt
@staff_member_required
def banco_ajax(request):
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'Método no permitido'})
    action = request.POST.get('action', '')
    if action == 'insert':
        Banco.objects.create(
            nom_banque=request.POST.get('nom_banque', ''),
            compte_banque=request.POST.get('compte_banque', ''),
            address_banque=request.POST.get('address_banque', ''),
            distrit_banque=request.POST.get('distrit_banque', ''),
            ville_banque=request.POST.get('ville_banque', ''),
        )
        return JsonResponse({'ok': True})
    if request.POST.get('edit_id'):
        b = get_object_or_404(Banco, pk=request.POST.get('edit_id'))
        return JsonResponse({'id': b.pk, 'nom_banque': b.nom_banque, 'compte_banque': b.compte_banque,
                             'address_banque': b.address_banque, 'distrit_banque': b.distrit_banque,
                             'ville_banque': b.ville_banque})
    if action == 'update':
        b = get_object_or_404(Banco, pk=request.POST.get('id_banque'))
        b.nom_banque = request.POST.get('nom_banque', b.nom_banque)
        b.compte_banque = request.POST.get('compte_banque', b.compte_banque)
        b.address_banque = request.POST.get('address_banque', b.address_banque)
        b.distrit_banque = request.POST.get('distrit_banque', b.distrit_banque)
        b.ville_banque = request.POST.get('ville_banque', b.ville_banque)
        b.save()
        return JsonResponse({'ok': True})
    if request.POST.get('del_id'):
        get_object_or_404(Banco, pk=request.POST.get('del_id')).delete()
        return JsonResponse({'ok': True})
    return JsonResponse({'ok': False, 'error': 'Acción no reconocida'})


# --- AJAX Benificiario (otro_benificiaire.php) ---

@csrf_exempt
@staff_member_required
def benificiario_ajax(request):
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'Método no permitido'})
    action = request.POST.get('action', '')
    if action == 'insert':
        Benificiario.objects.create(
            nom_benificiaire=request.POST.get('nom_benificiaire', ''),
            banque_benificiaire=request.POST.get('banque_benificiaire', ''),
            compte_benificiaire=request.POST.get('compte_benificiaire', ''),
        )
        return JsonResponse({'ok': True})
    if request.POST.get('edit_id'):
        b = get_object_or_404(Benificiario, pk=request.POST.get('edit_id'))
        return JsonResponse({'id': b.pk, 'nom_benificiaire': b.nom_benificiaire,
                             'banque_benificiaire': b.banque_benificiaire,
                             'compte_benificiaire': b.compte_benificiaire})
    if action == 'update':
        b = get_object_or_404(Benificiario, pk=request.POST.get('id_benificiaire'))
        b.nom_benificiaire = request.POST.get('nom_benificiaire', b.nom_benificiaire)
        b.banque_benificiaire = request.POST.get('banque_benificiaire', b.banque_benificiaire)
        b.compte_benificiaire = request.POST.get('compte_benificiaire', b.compte_benificiaire)
        b.save()
        return JsonResponse({'ok': True})
    if request.POST.get('del_id'):
        get_object_or_404(Benificiario, pk=request.POST.get('del_id')).delete()
        return JsonResponse({'ok': True})
    return JsonResponse({'ok': False, 'error': 'Acción no reconocida'})
