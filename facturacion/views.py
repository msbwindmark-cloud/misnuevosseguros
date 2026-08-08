import io
from decimal import Decimal
from datetime import datetime
from django.shortcuts import render, get_object_or_404, redirect
from django.http import HttpResponse
from django.contrib import messages
from django.db.models import Sum
from polizas.models import Poliza
from clientes.models import Cliente

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm, cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.colors import HexColor, black, white, Color
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT, TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.platypus.flowables import HRFlowable
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


COMPANY_NAME = "PROYECTO SIS"
COMPANY_ADDRESS = "123 Rue Example, 75000 Paris"
COMPANY_PHONE = "Tel: +33 1 23 45 67 89"
COMPANY_EMAIL = "info@proyectosis.com"
COMPANY_SIRET = "SIRET: 123 456 789 00012"


def _build_pdf_response(buffer, filename):
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _header_table(doc_title="FACTURE", doc_number=""):
    data = [
        [Paragraph(f"<b>{COMPANY_NAME}</b>", ParagraphStyle('hdr', fontSize=16, textColor=HexColor('#0d6efd'))),
         Paragraph(f"<b>{doc_title}</b><br/>{doc_number}", ParagraphStyle('hdrr', fontSize=14, alignment=TA_RIGHT))],
        [Paragraph(COMPANY_ADDRESS, ParagraphStyle('addr', fontSize=8)),
         Paragraph("", ParagraphStyle('addr'))],
        [Paragraph(COMPANY_PHONE, ParagraphStyle('addr', fontSize=8)),
         Paragraph("", ParagraphStyle('addr'))],
        [Paragraph(COMPANY_EMAIL, ParagraphStyle('addr', fontSize=8)),
         Paragraph("", ParagraphStyle('addr'))],
    ]
    t = Table(data, colWidths=[A4[0]/2-40, A4[0]/2-40])
    t.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('SPAN', (0, 1), (0, 3)),
        ('LINEBELOW', (0, 0), (-1, 0), 1, HexColor('#0d6efd')),
    ]))
    return t


def _client_info_table(poliza):
    data = [
        [Paragraph("<b>CLIENT</b>", ParagraphStyle('clbl', fontSize=10))],
        [Paragraph(f"{poliza.assure} {poliza.prenom}", ParagraphStyle('clv', fontSize=10))],
        [Paragraph(f"Code client: {poliza.code_client}", ParagraphStyle('clv', fontSize=10))],
    ]
    if poliza.address_client:
        data.append([Paragraph(poliza.address_client.replace('\n', '<br/>'), ParagraphStyle('clv', fontSize=10))])
    if poliza.telefono:
        data.append([Paragraph(f"Tel: {poliza.telefono}", ParagraphStyle('clv', fontSize=10))])
    if poliza.email:
        data.append([Paragraph(f"Email: {poliza.email}", ParagraphStyle('clv', fontSize=10))])
    t = Table(data, colWidths=[A4[0]-80])
    t.setStyle(TableStyle([
        ('BOX', (0, 0), (-1, -1), 0.5, HexColor('#cccccc')),
        ('BACKGROUND', (0, 0), (0, 0), HexColor('#f0f0f0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
    ]))
    return t


def _invoice_items_table(poliza, show_payment=True):
    data = [
        [Paragraph("<b>Description</b>", ParagraphStyle('th', fontSize=9)),
         Paragraph("<b>Montant</b>", ParagraphStyle('th', fontSize=9, alignment=TA_RIGHT))],
        [Paragraph(f"Police: {poliza.police}", ParagraphStyle('td', fontSize=9)),
         Paragraph("", ParagraphStyle('td'))],
        [Paragraph(f"Attestation: {poliza.attestation}", ParagraphStyle('td', fontSize=9)),
         Paragraph("", ParagraphStyle('td'))],
        [Paragraph(f"Type: {poliza.letype}", ParagraphStyle('td', fontSize=9)),
         Paragraph("", ParagraphStyle('td'))],
        [Paragraph(f"Produit: {poliza.produit}", ParagraphStyle('td', fontSize=9)),
         Paragraph("", ParagraphStyle('td'))],
        [Paragraph(f"Matricule: {poliza.matricule}", ParagraphStyle('td', fontSize=9)),
         Paragraph("", ParagraphStyle('td'))],
        [Paragraph(f"Du: {poliza.du.strftime('%d/%m/%Y') if poliza.du else 'N/A'}  Au: {poliza.au.strftime('%d/%m/%Y') if poliza.au else 'N/A'}", ParagraphStyle('td', fontSize=9)),
         Paragraph("", ParagraphStyle('td'))],
    ]
    if show_payment:
        data.append([Paragraph("<b>Prime totale</b>", ParagraphStyle('tdb', fontSize=10)),
                     Paragraph(f"<b>{poliza.totale} MAD</b>", ParagraphStyle('tdb', fontSize=10, alignment=TA_RIGHT))])
        data.append([Paragraph("Efectivo (Espèce)", ParagraphStyle('td', fontSize=9)),
                     Paragraph(f"{poliza.espece} MAD", ParagraphStyle('td', fontSize=9, alignment=TA_RIGHT))])
        data.append([Paragraph("Cheque", ParagraphStyle('td', fontSize=9)),
                     Paragraph(f"{poliza.cheque} MAD", ParagraphStyle('td', fontSize=9, alignment=TA_RIGHT))])
        data.append([Paragraph("Virement (Transferencia)", ParagraphStyle('td', fontSize=9)),
                     Paragraph(f"{poliza.virement} MAD", ParagraphStyle('td', fontSize=9, alignment=TA_RIGHT))])
        if poliza.reste > 0:
            data.append([Paragraph("<b>Reste à payer</b>", ParagraphStyle('tdb', fontSize=10, textColor=HexColor('#dc3545'))),
                         Paragraph(f"<b>{poliza.reste} MAD</b>", ParagraphStyle('tdb', fontSize=10, alignment=TA_RIGHT, textColor=HexColor('#dc3545')))])
        else:
            data.append([Paragraph("<b>Reste</b>", ParagraphStyle('tdb', fontSize=10)),
                         Paragraph(f"<b>0,00 MAD</b>", ParagraphStyle('tdb', fontSize=10, alignment=TA_RIGHT))])

    col_widths = [A4[0]-80-60, 60]
    t = Table(data, colWidths=col_widths)
    style_cmds = [
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, HexColor('#cccccc')),
        ('BACKGROUND', (0, 0), (-1, 0), HexColor('#0d6efd')),
        ('TEXTCOLOR', (0, 0), (-1, 0), white),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
    ]
    t.setStyle(TableStyle(style_cmds))
    return t


def _terms_section():
    data = [
        [Paragraph("<b>Conditions de paiement:</b>", ParagraphStyle('term', fontSize=8)),
         Paragraph("Paiement à réception. Merci de votre confiance.", ParagraphStyle('term', fontSize=8))],
    ]
    t = Table(data, colWidths=[A4[0]-80])
    t.setStyle(TableStyle([
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    return t


def _footer_section():
    return Paragraph("<i>Document généré par Proyecto SIS - Merci de votre confiance</i>",
                     ParagraphStyle('foot', fontSize=7, alignment=TA_CENTER, textColor=HexColor('#999999')))


def _quittance_pdf(poliza, doc_type="FACTURE"):
    from reportlab.pdfgen import canvas as canvasmod
    fecha_str = datetime.now().strftime('%d/%m/%Y')
    prefix = "F" if doc_type == "FACTURE" else "R"

    buf = io.BytesIO()
    c = canvasmod.Canvas(buf, pagesize=A4)
    c.setTitle(doc_type)

    y = A4[1] - 35*mm

    c.setFont('Helvetica', 12)
    c.drawString(60, y, f'Rabat, le :  {fecha_str}')
    y -= 12*mm
    c.setFont('Helvetica-Bold', 12)
    c.drawString(60, y, f'{poliza.assure} {poliza.prenom}')
    y -= 10*mm

    if poliza.address_client:
        c.setFillColor(HexColor('#e6e6e6'))
        addr_lines = [poliza.address_client[i:i+60] for i in range(0, len(poliza.address_client), 60)]
        addr_h = len(addr_lines) * 5*mm
        c.rect(60, y - addr_h + 3*mm, 120*mm, addr_h + 2*mm, fill=True, stroke=False)
        c.setFillColor(black)
        c.setFont('Helvetica', 9)
        for line in addr_lines:
            c.drawString(62, y, line)
            y -= 5*mm
        y -= 4*mm

    y -= 8*mm
    c.setFont('Helvetica-Bold', 14)
    c.drawCentredString(A4[0]/2, y, doc_type)
    tw = c.stringWidth(doc_type, 'Helvetica-Bold', 14)
    c.line(A4[0]/2 - tw/2, y - 2, A4[0]/2 + tw/2, y - 2)
    y -= 18*mm

    col_widths = [28*mm, 32*mm, 35*mm, 22*mm, 18*mm, 14*mm, 22*mm, 22*mm]
    headers = ['N\u00b0 Police', 'P\u00e9riode', 'D\u00e9signation', 'Prime Nette', 'Taxes', 'T.P', 'ACCESOIRES', 'Total TTC']
    total_w = sum(col_widths)
    x_start = (A4[0] - total_w) / 2

    c.setFont('Helvetica-Bold', 8)
    c.setFillColor(HexColor('#0D6EFD'))
    c.rect(x_start, y, total_w, 8*mm, fill=True, stroke=False)
    c.setFillColor(white)
    x = x_start
    for i, h in enumerate(headers):
        c.drawCentredString(x + col_widths[i]/2, y + 2.5*mm, h)
        x += col_widths[i]
    c.setFillColor(black)
    y -= 1*mm

    row_strs = [
        poliza.police or '',
        f"{poliza.du.strftime('%d/%m/%Y') if poliza.du else ''} AU {poliza.au.strftime('%d/%m/%Y') if poliza.au else ''}",
        poliza.designation or '',
        str(poliza.prime_net or 0),
        str(poliza.taxes or 0),
        str(poliza.tp or 0),
        str(poliza.accesoires or 0),
        str(poliza.totale or 0),
    ]
    max_lines = 1
    for val in row_strs:
        n = max(1, -(-len(val) // 20))
        max_lines = max(max_lines, n)
    row_h = max(12*mm, max_lines * 5*mm + 6*mm)

    c.setFont('Helvetica', 8)
    c.setStrokeColor(HexColor('#cccccc'))
    x = x_start
    for i, val in enumerate(row_strs):
        c.rect(x, y - row_h, col_widths[i], row_h, fill=False, stroke=True)
        wrapped = [val[j:j+20] for j in range(0, len(val), 20)]
        ty = y - 4*mm
        for wl in wrapped:
            c.drawCentredString(x + col_widths[i]/2, ty, wl)
            ty -= 4.5*mm
        x += col_widths[i]
    y -= row_h

    y -= 10*mm
    totals = [
        ("Total Primes Nettes:", str(poliza.prime_net or 0)),
        ("Total Taxes & T.P.:", f"{float(poliza.taxes or 0) + float(poliza.tp or 0)}"),
        ("Total Accessoires:", str(poliza.accesoires or 0)),
        ("Total TTC:", str(poliza.totale or 0)),
    ]
    c.setFont('Helvetica-Bold', 10)
    for label, val in totals:
        c.setStrokeColor(HexColor('#cccccc'))
        c.rect(x_start + total_w - 130*mm, y, 90*mm, 7*mm, fill=False, stroke=True)
        c.drawString(x_start + total_w - 128*mm, y + 2*mm, label)
        c.rect(x_start + total_w - 40*mm, y, 40*mm, 7*mm, fill=False, stroke=True)
        c.drawRightString(x_start + total_w - 3*mm, y + 2*mm, val)
        y -= 7*mm

    c.save()
    buf.seek(0)
    return buf


# ---------- INVOICE VIEWS ----------

def invoice_by_receipt(request, pk):
    poliza = get_object_or_404(Poliza, recu=pk)
    buf = _quittance_pdf(poliza, "FACTURE")
    now = datetime.now()
    ts = f"{now.day:02d}-{now.month:02d}-{now.year}-{now.hour}-{now.minute:02d}-{now.second:02d}"
    return _build_pdf_response(buf, f'{poliza.assure}-F-{str(poliza.recu).zfill(5)}-{ts}.pdf')


def _client_pdf(cliente, polizas, prefix, title):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm, leftMargin=2*cm, rightMargin=2*cm)
    elements = []
    fecha_hoy = datetime.now().strftime('%d/%m/%Y')
    year_str = datetime.now().strftime('%Y')
    ps_normal = ParagraphStyle('n', fontSize=8, leading=10)
    ps_bold = ParagraphStyle('b', fontSize=8, leading=10, fontName='Helvetica-Bold')
    ps_th = ParagraphStyle('th', fontSize=8, leading=10, fontName='Helvetica-Bold', textColor=white)
    ps_title = ParagraphStyle('title', fontSize=14, leading=16, fontName='Helvetica-Bold', alignment=TA_CENTER, textColor=HexColor('#DC3545'))
    ps_addr = ParagraphStyle('addr', fontSize=9, leading=11)

    elements.append(Spacer(1, 30))
    elements.append(Paragraph(f"Rabat, le : {fecha_hoy}", ps_normal))
    elements.append(Spacer(1, 10))
    elements.append(Paragraph(f"<b>{cliente.nom_client}</b>", ps_bold))
    elements.append(Spacer(1, 4))
    if cliente.telephone_client:
        elements.append(Paragraph(f"Tel: {cliente.telephone_client}", ps_addr))
    first_poliza = polizas.first()
    if first_poliza and first_poliza.address_client:
        elements.append(Paragraph(first_poliza.address_client, ps_addr))
        elements.append(Spacer(1, 10))

    recus_str = '-'.join(str(p.recu) for p in polizas)
    id_client = cliente.code_client
    elements.append(Paragraph(f"<u>{prefix}  N\u00ba: {prefix}-{id_client}__{recus_str}_{year_str}</u>", ps_title))
    elements.append(Spacer(1, 15))

    header = [
        Paragraph("<b>N\u00b0 Police</b>", ps_th),
        Paragraph("<b>P\u00e9riode</b>", ps_th),
        Paragraph("<b>D\u00e9signation</b>", ps_th),
        Paragraph("<b>Prime Nette</b>", ps_th),
        Paragraph("<b>Taxes</b>", ps_th),
        Paragraph("<b>T.P</b>", ps_th),
        Paragraph("<b>ACCESOIRES</b>", ps_th),
        Paragraph("<b>Total TTC</b>", ps_th),
    ]
    data = [header]
    tot_pn = tot_tax = tot_tp = tot_acc = tot_ttc = 0
    for p in polizas:
        pn = float(p.prime_net or 0)
        tax = float(p.taxes or 0)
        tp = float(p.tp or 0)
        acc = float(p.accesoires or 0)
        tot = float(p.totale or 0)
        tot_pn += pn
        tot_tax += tax
        tot_tp += tp
        tot_acc += acc
        tot_ttc += tot
        period = f"{p.du.strftime('%d/%m/%Y') if p.du else ''} - {p.au.strftime('%d/%m/%Y') if p.au else ''}"
        data.append([
            Paragraph(p.police, ps_normal), Paragraph(period, ps_normal), Paragraph(p.designation, ps_normal),
            Paragraph(str(pn), ps_normal), Paragraph(str(tax), ps_normal), Paragraph(str(tp), ps_normal),
            Paragraph(str(acc), ps_normal), Paragraph(str(tot), ps_normal),
        ])
    data.append([Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("<b>Total</b>", ps_bold),
                 Paragraph(f"<b>{tot_pn}</b>", ps_bold), Paragraph(f"<b>{tot_tax}</b>", ps_bold),
                 Paragraph(f"<b>{tot_tp}</b>", ps_bold), Paragraph(f"<b>{tot_acc}</b>", ps_bold),
                 Paragraph(f"<b>{tot_ttc}</b>", ps_bold)])
    data.append([Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("<b>Total Primes Nettes:</b>", ps_bold),
                 Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("", ps_bold),
                 Paragraph(f"<b>{tot_pn}</b>", ps_bold)])
    data.append([Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("<b>Total Taxes & T.P.:</b>", ps_bold),
                 Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("", ps_bold),
                 Paragraph(f"<b>{tot_tax + tot_tp}</b>", ps_bold)])
    data.append([Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("<b>Total Accessoires:</b>", ps_bold),
                 Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("", ps_bold),
                 Paragraph(f"<b>{tot_acc}</b>", ps_bold)])
    data.append([Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("<b>Total TTC:</b>", ps_bold),
                 Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("", ps_bold), Paragraph("", ps_bold),
                 Paragraph(f"<b>{tot_ttc}</b>", ps_bold)])

    avail = A4[0] - 4*cm
    cw = [avail*0.17, avail*0.2, avail*0.18, avail*0.11, avail*0.09, avail*0.08, avail*0.1, avail*0.11]
    t = Table(data, colWidths=cw)
    bg_grey = HexColor('#E8E8E8')
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), HexColor('#0D6EFD')),
        ('TEXTCOLOR', (0, 0), (-1, 0), white),
        ('BACKGROUND', (0, -6), (-1, -6), bg_grey),
        ('BACKGROUND', (0, -5), (-1, -5), bg_grey),
        ('BACKGROUND', (0, -4), (-1, -4), bg_grey),
        ('BACKGROUND', (0, -3), (-1, -3), bg_grey),
        ('BACKGROUND', (0, -2), (-1, -2), bg_grey),
        ('BACKGROUND', (0, -1), (-1, -1), HexColor('#F0F0F0')),
        ('GRID', (0, 0), (-1, -1), 0.5, HexColor('#CCCCCC')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(t)
    elements.append(_footer_section())
    doc.build(elements)
    return buf


def invoice_by_client(request, client_id):
    cliente = get_object_or_404(Cliente, code_client=client_id)
    polizas = Poliza.objects.filter(code_client=cliente.code_client).order_by('-recu')
    if not polizas.exists():
        return render(request, 'facturacion/sin_polizas.html', {'modal': 'factures'})
    buf = _client_pdf(cliente, polizas, "F", "FACTURE")
    return _build_pdf_response(buf, f'{client_id}-F-global-{datetime.now().strftime("%d-%m-%Y-%H-%M-%S")}.pdf')


def _quittance_excel(poliza, prefix="F"):
    wb = openpyxl.Workbook()
    ws = wb.active
    fecha_str = datetime.now().strftime('%d/%m/%Y')
    year_str = datetime.now().strftime('%Y')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    label_font = Font(name='Calibri', bold=True, size=14)
    normal_font = Font(name='Calibri', size=14)
    bold_font = Font(name='Calibri', bold=True, size=14)
    title_font = Font(name='Calibri', bold=True, size=14)
    center = Alignment(horizontal='center', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')

    for r in range(1, 14):
        ws.row_dimensions[r].height = 2

    ws.cell(row=14, column=3, value=f'Rabat, le : {fecha_str}').font = normal_font
    ws.cell(row=16, column=3, value=f'{poliza.assure} {poliza.prenom}').font = bold_font
    ws.cell(row=18, column=3, value=poliza.address_client or '').font = normal_font
    ws.cell(row=20, column=3, value=f'{prefix}  N\u00ba: {prefix}-{str(poliza.recu).zfill(5)}-{year_str}').font = title_font

    headers = ['N\u00b0 Police', 'P\u00e9riode', 'D\u00e9signation', 'Prime Nette', 'Taxes', 'T.P', 'ACCESOIRES', 'Total TTC']
    header_row = 22
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = Font(name='Calibri', bold=True, size=14)
        c.border = thin_border
        c.alignment = center

    data_row = header_row + 1
    ws.cell(row=data_row, column=1, value=poliza.police).border = thin_border
    ws.cell(row=data_row, column=2, value=f"{poliza.du.strftime('%d/%m/%Y') if poliza.du else ''} - {poliza.au.strftime('%d/%m/%Y') if poliza.au else ''}").border = thin_border
    ws.cell(row=data_row, column=3, value=poliza.designation).border = thin_border
    for ci, val in enumerate([poliza.prime_net, poliza.taxes, poliza.tp, poliza.accesoires, poliza.totale], 4):
        cell = ws.cell(row=data_row, column=ci, value=str(float(val)).replace('.', ','))
        cell.border = thin_border
        cell.alignment = right_align

    tot_row = data_row + 2
    for label, val in [('Total Primes', poliza.prime_net), ('Total Taxes', poliza.taxes),
                       ('Total Timbres & Accessoires', poliza.accesoires), ('Total TTC', poliza.totale)]:
        ws.cell(row=tot_row, column=3, value=label).font = bold_font
        ws.cell(row=tot_row, column=3).border = thin_border
        cell = ws.cell(row=tot_row, column=4, value=str(float(val)).replace('.', ','))
        cell.font = bold_font
        cell.border = thin_border
        cell.alignment = right_align
        tot_row += 1

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 8
    ws.column_dimensions['G'].width = 18
    ws.column_dimensions['H'].width = 15

    now = datetime.now()
    ts = f"{now.day:02d}-{now.month:02d}-{now.year}-{now.hour}-{now.minute:02d}-{now.second:02d}"
    return wb, f'{poliza.assure}-{prefix}-{str(poliza.recu).zfill(5)}-{ts}.xls'


def invoice_by_receipt_excel(request, pk):
    poliza = get_object_or_404(Poliza, recu=pk)
    wb, filename = _quittance_excel(poliza, "F")
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


# ---------- STATEMENT VIEWS ----------

def statement_by_receipt(request, pk):
    poliza = get_object_or_404(Poliza, recu=pk)
    buf = _quittance_pdf(poliza, "RELEVE")
    now = datetime.now()
    ts = f"{now.day:02d}-{now.month:02d}-{now.year}-{now.hour}-{now.minute:02d}-{now.second:02d}"
    return _build_pdf_response(buf, f'{poliza.assure}-R-{str(poliza.recu).zfill(5)}-{ts}.pdf')


def statement_by_client(request, client_id):
    cliente = get_object_or_404(Cliente, code_client=client_id)
    polizas = Poliza.objects.filter(code_client=cliente.code_client).order_by('-recu')
    if not polizas.exists():
        return render(request, 'facturacion/sin_polizas.html', {'modal': 'releves'})
    buf = _client_pdf(cliente, polizas, "R", "RELEVÉ")
    return _build_pdf_response(buf, f'{client_id}-R-global-{datetime.now().strftime("%d-%m-%Y-%H-%M-%S")}.pdf')


def statement_by_receipt_excel(request, pk):
    poliza = get_object_or_404(Poliza, recu=pk)
    wb, filename = _quittance_excel(poliza, "R")
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


# ---------- BALANCE PENDIENTE ----------

def balance_pendiente_list(request):
    polizas = Poliza.objects.filter(reste__gt=0).order_by('-recu')
    total_general = polizas.aggregate(
        tot=Sum('totale'), res=Sum('reste')
    )

    if 'export_excel' in request.GET:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Balance Pendiente"
        headers = ['Recibo', 'Asegurado', 'Póliza', 'Total', 'Pagado', 'Reste', 'Teléfono']
        header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
        header_font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
        for ri, p in enumerate(polizas, 2):
            ws.cell(row=ri, column=1, value=p.recu).border = thin_border
            ws.cell(row=ri, column=2, value=str(p.assure)).border = thin_border
            ws.cell(row=ri, column=3, value=p.police).border = thin_border
            cell_t = ws.cell(row=ri, column=4, value=float(p.totale))
            cell_t.number_format = '#,##0.00'
            cell_t.border = thin_border
            pagado = float(p.totale - p.reste)
            cell_p = ws.cell(row=ri, column=5, value=pagado)
            cell_p.number_format = '#,##0.00'
            cell_p.border = thin_border
            cell_r = ws.cell(row=ri, column=6, value=float(p.reste))
            cell_r.number_format = '#,##0.00'
            cell_r.border = thin_border
            ws.cell(row=ri, column=7, value=p.telefono).border = thin_border
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 18
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="balance_pendiente.xlsx"'
        wb.save(response)
        return response

    pagado_total = (total_general['tot'] or 0) - (total_general['res'] or 0)
    for p in polizas:
        p.pagado = p.totale - p.reste

    return render(request, 'facturacion/balance_pendiente.html', {
        'polizas': polizas,
        'total_general': total_general,
        'pagado_total': pagado_total,
    })


# ---------- CONSULTA VENCIMIENTO ----------

def consulta_vencimiento(request):
    resultados = None
    fecha_desde = request.GET.get('fecha_desde', '')
    fecha_hasta = request.GET.get('fecha_hasta', '')

    if fecha_desde and fecha_hasta:
        resultados = Poliza.objects.filter(au__gte=fecha_desde, au__lte=fecha_hasta).order_by('au')

    if 'export_excel' in request.GET and resultados is not None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Vencimientos"
        headers = ['Recibo', 'Asegurado', 'Póliza', 'Desde', 'Hasta', 'Teléfono', 'Email']
        header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
        header_font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
        for ri, p in enumerate(resultados, 2):
            ws.cell(row=ri, column=1, value=p.recu).border = thin_border
            ws.cell(row=ri, column=2, value=str(p.assure)).border = thin_border
            ws.cell(row=ri, column=3, value=p.police).border = thin_border
            ws.cell(row=ri, column=4, value=p.du.strftime('%d/%m/%Y') if p.du else '').border = thin_border
            ws.cell(row=ri, column=5, value=p.au.strftime('%d/%m/%Y') if p.au else '').border = thin_border
            ws.cell(row=ri, column=6, value=p.telefono).border = thin_border
            ws.cell(row=ri, column=7, value=p.email).border = thin_border
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 14
        ws.column_dimensions['E'].width = 14
        ws.column_dimensions['F'].width = 18
        ws.column_dimensions['G'].width = 30
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="consulta_vencimiento.xlsx"'
        wb.save(response)
        return response

    return render(request, 'facturacion/consulta_vencimiento.html', {
        'resultados': resultados,
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
    })


def invoice_list_by_client(request, client_id):
    polizas = Poliza.objects.filter(code_client=client_id).order_by('-recu')
    cliente = Cliente.objects.filter(code_client=client_id).first()
    cliente_nom = cliente.nom_client if cliente else str(client_id)
    cliente_adresse = ''
    primer_recu = ''
    if polizas.exists():
        primer_recu = polizas.first().recu
        cliente_adresse = polizas.first().address_client or ''

    if 'export_excel' in request.GET:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f'Facturas {client_id}'
        fecha_str = datetime.now().strftime('%d/%m/%Y')
        facture_num = f"F-{client_id}__{primer_recu}_{fecha_str}"
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        for r in range(1, 14):
            ws.row_dimensions[r].height = 2

        ws.merge_cells('A1:H1')
        ws['A1'] = f'Rabat, le : {fecha_str}'
        ws['A1'].font = Font(name='Calibri', size=14)
        ws.merge_cells('A3:H3')
        ws['A3'] = cliente_nom
        ws['A3'].font = Font(name='Calibri', bold=True, size=14)
        ws.merge_cells('A5:H5')
        ws['A5'] = cliente_adresse
        ws['A5'].font = Font(name='Calibri', size=14)
        ws.merge_cells('A7:H7')
        ws['A7'] = f'FACTURE  N\u00ba: {facture_num}'
        ws['A7'].font = Font(name='Calibri', bold=True, size=14)

        header_row = 9
        headers = ['N\u00b0 Police', 'P\u00e9riode', 'D\u00e9signation', 'Prime Nette', 'Taxes', 'T.P', 'Accessoires', 'Total TTC']
        header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
        header_font = Font(name='Calibri', bold=True, size=14, color='FFFFFF')
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=header_row, column=col, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal='center')
            c.border = thin_border

        data_row = header_row + 1
        tot_prime = tot_taxes = tot_tp = tot_acces = tot_ttc = 0
        for p in polizas:
            ws.cell(row=data_row, column=1, value=p.police).border = thin_border
            periode = f"{p.du.strftime('%d/%m/%Y') if p.du else ''} - {p.au.strftime('%d/%m/%Y') if p.au else ''}"
            ws.cell(row=data_row, column=2, value=periode).border = thin_border
            ws.cell(row=data_row, column=3, value=p.designation).border = thin_border
            ws.cell(row=data_row, column=4, value=str(float(p.prime_net)).replace('.', ',')).border = thin_border
            ws.cell(row=data_row, column=5, value=str(float(p.taxes)).replace('.', ',')).border = thin_border
            ws.cell(row=data_row, column=6, value=str(float(p.tp)).replace('.', ',')).border = thin_border
            ws.cell(row=data_row, column=7, value=str(float(p.accesoires)).replace('.', ',')).border = thin_border
            ws.cell(row=data_row, column=8, value=str(float(p.totale)).replace('.', ',')).border = thin_border
            tot_prime += float(p.prime_net)
            tot_taxes += float(p.taxes)
            tot_tp += float(p.tp)
            tot_acces += float(p.accesoires)
            tot_ttc += float(p.totale)
            data_row += 1

        ws.cell(row=data_row, column=3, value='Total Primes Nettes').font = Font(bold=True)
        ws.cell(row=data_row, column=4, value=str(tot_prime).replace('.', ',')).font = Font(bold=True)
        data_row += 1
        ws.cell(row=data_row, column=3, value='Total Taxes & T.P').font = Font(bold=True)
        ws.cell(row=data_row, column=4, value=str(tot_taxes + tot_tp).replace('.', ',')).font = Font(bold=True)
        data_row += 1
        ws.cell(row=data_row, column=3, value='Total Accessoires').font = Font(bold=True)
        ws.cell(row=data_row, column=4, value=str(tot_acces).replace('.', ',')).font = Font(bold=True)
        data_row += 1
        ws.cell(row=data_row, column=3, value='Total TTC').font = Font(bold=True)
        ws.cell(row=data_row, column=4, value=str(tot_ttc).replace('.', ',')).font = Font(bold=True)

        for col_letter in ['A','B','C','D','E','F','G','H']:
            ws.column_dimensions[col_letter].width = 18
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{client_id}-F-{str(primer_recu).zfill(5)}-{datetime.now().strftime("%d-%m-%Y-%H-%M-%S")}.xls"'
        wb.save(response)
        return response

    return render(request, 'facturacion/invoice_list_client.html', {
        'polizas': polizas,
        'cliente_nom': cliente_nom,
        'cliente_code': client_id,
        'cliente_adresse': cliente_adresse,
        'primer_recu': primer_recu,
        'fecha_hoy': datetime.now().strftime('%d/%m/%Y'),
    })


def statement_list_by_client(request, client_id):
    polizas = Poliza.objects.filter(code_client=client_id).order_by('-recu')
    cliente = Cliente.objects.filter(code_client=client_id).first()
    cliente_nom = cliente.nom_client if cliente else str(client_id)
    cliente_adresse = ''
    primer_recu = ''
    if polizas.exists():
        primer_recu = polizas.first().recu
        cliente_adresse = polizas.first().address_client or ''

    if 'export_excel' in request.GET:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f'Relevés {client_id}'
        fecha_str = datetime.now().strftime('%d/%m/%Y')
        releve_num = f"R-{client_id}__{primer_recu}_{fecha_str}"
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        for r in range(1, 14):
            ws.row_dimensions[r].height = 2

        ws.merge_cells('A1:H1')
        ws['A1'] = f'Rabat, le : {fecha_str}'
        ws['A1'].font = Font(name='Calibri', size=14)
        ws.merge_cells('A3:H3')
        ws['A3'] = cliente_nom
        ws['A3'].font = Font(name='Calibri', bold=True, size=14)
        ws.merge_cells('A5:H5')
        ws['A5'] = cliente_adresse
        ws['A5'].font = Font(name='Calibri', size=14)
        ws.merge_cells('A7:H7')
        ws['A7'] = f'RELEVÉ  N\u00ba: {releve_num}'
        ws['A7'].font = Font(name='Calibri', bold=True, size=14)

        header_row = 9
        headers = ['N\u00b0 Police', 'P\u00e9riode', 'D\u00e9signation', 'Prime Nette', 'Taxes', 'T.P', 'Accessoires', 'Total TTC']
        header_fill = PatternFill(start_color='198754', end_color='198754', fill_type='solid')
        header_font = Font(name='Calibri', bold=True, size=14, color='FFFFFF')
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=header_row, column=col, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal='center')
            c.border = thin_border

        data_row = header_row + 1
        tot_prime = tot_taxes = tot_tp = tot_acces = tot_ttc = 0
        for p in polizas:
            ws.cell(row=data_row, column=1, value=p.police).border = thin_border
            periode = f"{p.du.strftime('%d/%m/%Y') if p.du else ''} - {p.au.strftime('%d/%m/%Y') if p.au else ''}"
            ws.cell(row=data_row, column=2, value=periode).border = thin_border
            ws.cell(row=data_row, column=3, value=p.designation).border = thin_border
            ws.cell(row=data_row, column=4, value=str(float(p.prime_net)).replace('.', ',')).border = thin_border
            ws.cell(row=data_row, column=5, value=str(float(p.taxes)).replace('.', ',')).border = thin_border
            ws.cell(row=data_row, column=6, value=str(float(p.tp)).replace('.', ',')).border = thin_border
            ws.cell(row=data_row, column=7, value=str(float(p.accesoires)).replace('.', ',')).border = thin_border
            ws.cell(row=data_row, column=8, value=str(float(p.totale)).replace('.', ',')).border = thin_border
            tot_prime += float(p.prime_net)
            tot_taxes += float(p.taxes)
            tot_tp += float(p.tp)
            tot_acces += float(p.accesoires)
            tot_ttc += float(p.totale)
            data_row += 1

        ws.cell(row=data_row, column=3, value='Total Primes Nettes').font = Font(bold=True)
        ws.cell(row=data_row, column=4, value=str(tot_prime).replace('.', ',')).font = Font(bold=True)
        data_row += 1
        ws.cell(row=data_row, column=3, value='Total Taxes & T.P').font = Font(bold=True)
        ws.cell(row=data_row, column=4, value=str(tot_taxes + tot_tp).replace('.', ',')).font = Font(bold=True)
        data_row += 1
        ws.cell(row=data_row, column=3, value='Total Accessoires').font = Font(bold=True)
        ws.cell(row=data_row, column=4, value=str(tot_acces).replace('.', ',')).font = Font(bold=True)
        data_row += 1
        ws.cell(row=data_row, column=3, value='Total TTC').font = Font(bold=True)
        ws.cell(row=data_row, column=4, value=str(tot_ttc).replace('.', ',')).font = Font(bold=True)

        for col_letter in ['A','B','C','D','E','F','G','H']:
            ws.column_dimensions[col_letter].width = 18
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{client_id}-R-{str(primer_recu).zfill(5)}-{datetime.now().strftime("%d-%m-%Y-%H-%M-%S")}.xls"'
        wb.save(response)
        return response

    return render(request, 'facturacion/statement_list_client.html', {
        'polizas': polizas,
        'cliente_nom': cliente_nom,
        'cliente_code': client_id,
        'cliente_adresse': cliente_adresse,
        'primer_recu': primer_recu,
        'fecha_hoy': datetime.now().strftime('%d/%m/%Y'),
    })
