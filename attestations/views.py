from django.urls import reverse_lazy
from django.http import JsonResponse, HttpResponse
from django.views.generic import ListView, UpdateView
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.admin.views.decorators import staff_member_required
from django.utils.decorators import method_decorator
from .models import Attestation, TypeAttestation
from .forms import AttestationForm
import openpyxl


class AttestationListView(ListView):
    model = Attestation
    template_name = 'attestations/lista.html'
    context_object_name = 'attestations'
    paginate_by = 100


@method_decorator(staff_member_required, name='dispatch')
class AttestationUpdateView(UpdateView):
    model = Attestation
    form_class = AttestationForm
    template_name = 'attestations/form.html'
    success_url = reverse_lazy('attestation_list')

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['tipos'] = TypeAttestation.objects.all()
        return ctx


@staff_member_required
def import_attestations_excel(request):
    if request.method == 'POST' and request.FILES.get('file'):
        try:
            wb = openpyxl.load_workbook(request.FILES['file'], data_only=True)
        except Exception as e:
            messages.error(request, f'Error al leer el archivo: {e}')
            return redirect('attestation_list')

        tipos_por_nombre = {t.type_attestation.strip().upper(): t for t in TypeAttestation.objects.all()}
        total = 0
        duplicados = 0
        importados_por_tipo = {}

        for idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            tipo = tipos_por_nombre.get(sheet_name.strip().upper())
            if not tipo:
                tipo = TypeAttestation.objects.filter(id=idx + 1).first()
            if not tipo:
                continue

            contador_tipo = 0
            for row in ws.iter_rows(values_only=True):
                if len(row) < 2:
                    continue
                valor = row[1]
                if valor is None or str(valor).strip() == '':
                    continue
                numero = str(valor).strip()
                if not Attestation.objects.filter(numero_attestation=numero, id_type_attestation=tipo).exists():
                    Attestation.objects.create(
                        id_type_attestation=tipo,
                        numero_attestation=numero,
                        seleccionado=True,
                    )
                    total += 1
                    contador_tipo += 1
                else:
                    duplicados += 1
            if contador_tipo:
                importados_por_tipo[tipo.type_attestation] = contador_tipo

        msg = f'{total} attestations importadas.'
        if importados_por_tipo:
            detalle = ', '.join(f'{k}: {v}' for k, v in importados_por_tipo.items())
            msg += f' ({detalle}).'
        if duplicados:
            msg += f' {duplicados} duplicados omitidos.'
        messages.success(request, msg)
        return redirect('attestation_list')

    tipos = TypeAttestation.objects.all()
    return render(request, 'attestations/import.html', {'tipos': tipos})


@staff_member_required
def attestation_delete(request):
    if request.method == 'POST':
        raw = request.POST.get('ids') or request.POST.get('ids[]')
        if not raw:
            ids_list = request.POST.getlist('ids[]')
        else:
            ids_list = [i.strip() for i in str(raw).split(',') if i.strip()]
        if ids_list:
            cantidad = Attestation.objects.filter(id__in=ids_list).delete()[0]
            messages.success(request, f'{cantidad} atestación(es) eliminada(s).')
        else:
            messages.error(request, 'No se seleccionaron registros.')
    return redirect('attestation_list')
