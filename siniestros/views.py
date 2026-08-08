import json, os
from datetime import datetime
from django.conf import settings
from django.db.models import Count
from django.urls import reverse_lazy
from django.http import JsonResponse, HttpResponse, FileResponse, HttpResponseNotFound
from django.views.generic import ListView, View
from django.contrib.admin.views.decorators import staff_member_required
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from .models import Sinistro, DocumentoSinistro, Formule, ReferenceReglement, DocumentDeBase
from polizas.models import Poliza
from expertos.models import Experto
from documentos.models import TipoDocumento, DocumentoUsuario


class SinistroListView(ListView):
    model = Sinistro
    template_name = 'siniestros/lista.html'
    context_object_name = 'siniestros'
    paginate_by = 50

    def get_queryset(self):
        qs = super().get_queryset()
        if self.request.GET.get('con_docs') == '1':
            qs = qs.filter(documentos__isnull=False).distinct()
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ids = list(Sinistro.objects.values_list('id', flat=True))
        from documentos.models import DocumentoUsuario
        doc_ids = set(DocumentoUsuario.objects.filter(id_usuario__in=ids).values_list('id_usuario', flat=True).distinct())
        ctx['siniestros_con_docs'] = doc_ids
        doc_count_qs = DocumentoSinistro.objects.filter(sinistro_id__in=ids).values('sinistro_id').annotate(total=Count('id')).values_list('sinistro_id', 'total')
        ctx['sinistro_docs_count'] = {r: t for r, t in doc_count_qs}
        ctx['filtrando_docs'] = self.request.GET.get('con_docs') == '1'
        ctx['tipos_doc'] = TipoDocumento.objects.all()
        return ctx


class SinistroCreateView(View):
    template_name = 'siniestros/form.html'

    def get(self, request):
        # Último código de siniestro para sugerencia
        ultimo = Sinistro.objects.order_by('-id').first()
        año = datetime.now().strftime('%y')
        if ultimo and ultimo.code_sinistre:
            try:
                num = int(ultimo.code_sinistre.split('-')[1].split('/')[0]) + 1
            except (IndexError, ValueError):
                num = 1
        else:
            num = 1
        sugerido = f'S-{num}/{año}'

        val = {k: '' for k in ['code_sinistre','ref_sinistre','cod_sig_correcto','completo','date_operation','date_sinistre','date_reception','police','assure','categorie_assurance','reference_assurance','benificiaire','reference_assistance','nom_formule','nom_expert','reference_reglement','document_de_base','montant','reference_payement','date_payement','id_recu']}
        val['code_sinistre'] = sugerido
        val['cod_sig_correcto'] = '0'
        return render(request, self.template_name, {
            'val': val,
            'sugerido': sugerido,
            'formules': Formule.objects.all(),
            'expertos': Experto.objects.all(),
            'referencias_reg': ReferenceReglement.objects.all(),
            'documentos_base': DocumentDeBase.objects.all(),
            'tipos_doc': TipoDocumento.objects.all(),
            'hoy': datetime.now().strftime('%Y-%m-%d'),
        })

    def post(self, request):
        data = request.POST
        files = request.FILES

        poliza = None
        id_recu = data.get('id_recu')
        if id_recu and id_recu.isdigit():
            poliza = Poliza.objects.filter(recu=int(id_recu)).first()

        sinistro = Sinistro.objects.create(
            code_sinistre=data.get('code_sinistre'),
            ref_sinistre=data.get('ref_sinistre'),
            date_operation=data.get('date_operation') or None,
            date_sinistre=data.get('date_sinistre') or None,
            date_reception=data.get('date_reception') or None,
            police=data.get('police', ''),
            nom_formule=data.get('nom_formule', ''),
            assure=data.get('assure', ''),
            categorie_assurance=data.get('categorie_assurance', ''),
            reference_assurance=data.get('reference_assurance', ''),
            benificiaire=data.get('benificiaire', ''),
            reference_assistance=data.get('reference_assistance', ''),
            nom_expert=data.get('nom_expert', ''),
            reference_reglement=data.get('reference_reglement', ''),
            document_de_base=data.get('document_de_base', ''),
            reference_payement=data.get('reference_payement', ''),
            montant=float(data.get('montant', 0)) if data.get('montant') else None,
            date_payement=data.get('date_payement') or None,
            cod_sig_correcto=data.get('cod_sig_correcto') == '1',
            completo=data.get('completo') == '1',
            poliza=poliza,
        )

        # --- Procesar documentos ---
        code_limpio = sinistro.code_sinistre.replace('/', '-') if sinistro.code_sinistre else f'SIN-{sinistro.id}'
        upload_dir = 'uploads_sinistres/'

        # 1. Documentos reutilizados de la póliza
        docs_reutilizados = data.getlist('docs_reutilizados[]')
        for id_doc in docs_reutilizados:
            doc_orig = DocumentoUsuario.objects.filter(id=id_doc).first()
            if doc_orig:
                DocumentoSinistro.objects.create(
                    code_sinistre=sinistro.code_sinistre,
                    nombre_archivo=doc_orig.nombre_documento_usuario,
                    tipo_documento=doc_orig.tipo_documento_id,
                    id_documento_usuario_origen=doc_orig.id,
                    sinistro=sinistro,
                )

        # 2. Subir nuevos documentos
        tipos_doc = data.getlist('tipo_doc_siniestro[]')
        for idx, id_tipo in enumerate(tipos_doc):
            if not id_tipo:
                continue
            file_key = 'archivos_siniestro_0' if idx == 0 else f'archivos_siniestro_{idx}'
            archivos = files.getlist(file_key)
            for i, f in enumerate(archivos):
                ext = f.name.split('.')[-1].lower() if '.' in f.name else ''
                nuevo_nombre = f'SIN_{code_limpio}_{datetime.now().timestamp()}_{i}.{ext}'
                doc = DocumentoSinistro(
                    code_sinistre=sinistro.code_sinistre,
                    nombre_archivo=nuevo_nombre,
                    tipo_documento=int(id_tipo) if id_tipo.isdigit() else None,
                    ruta_archivo=upload_dir + nuevo_nombre,
                    archivo=f,
                    sinistro=sinistro,
                )
                doc.save()

        messages.success(request, f'Siniestro {sinistro.code_sinistre} creado correctamente.')
        return redirect('sinistro_list')


def get_polizas_json(request):
    q = request.GET.get('q', '')
    polizas = Poliza.objects.all()
    if q:
        polizas = polizas.filter(police__icontains=q) | polizas.filter(assure__icontains=q)
    polizas = polizas.order_by('-recu')[:50]
    data = [{
        'recu': p.recu,
        'police': p.police,
        'assure': p.assure,
        'marque': p.marque,
        'matricule': p.matricule,
        'produit': p.produit,
        'code_client': p.code_client,
        'telefono': p.telefono,
        'email': p.email,
    } for p in polizas]
    return JsonResponse(data, safe=False)


def get_poliza_documentos(request, pk):
    docs = DocumentoUsuario.objects.filter(id_usuario=pk).select_related('tipo_documento')
    data = [{
        'id': d.id,
        'nombre': d.nombre_documento_usuario,
        'tipo': str(d.tipo_documento) if d.tipo_documento else '',
        'descargar': reverse_lazy('descargar_documento', kwargs={'pk': d.id}),
    } for d in docs]
    return JsonResponse(data, safe=False)


def sinistro_documentos_data(request, pk):
    sinistro = get_object_or_404(Sinistro, pk=pk)
    docs = DocumentoSinistro.objects.filter(sinistro=sinistro).order_by('-id')
    data = [{
        'id': d.id,
        'nombre': d.nombre_archivo or '',
        'tipo': str(d.tipo_documento) if d.tipo_documento else '—',
        'archivo': d.archivo.url if d.archivo else '',
        'descargar': reverse_lazy('descargar_documento_siniestro', kwargs={'pk': d.id}),
        'origen': d.id_documento_usuario_origen,
        'heredado': bool(d.id_documento_usuario_origen),
    } for d in docs]
    return JsonResponse({'documentos': data})


@staff_member_required
def sinistro_documentos_agregar(request, pk):
    from datetime import datetime as dt
    sinistro = get_object_or_404(Sinistro, pk=pk)
    if request.method == 'POST':
        data = request.POST
        files = request.FILES
        code_limpio = sinistro.code_sinistre.replace('/', '-') if sinistro.code_sinistre else f'SIN-{sinistro.id}'
        upload_dir = 'uploads_sinistres/'
        count = 0

        tipos_doc = data.getlist('tipo_doc_siniestro[]')
        for idx, id_tipo in enumerate(tipos_doc):
            if not id_tipo:
                continue
            file_key = f'archivos_siniestro_{idx}'
            archivos = files.getlist(file_key)
            for i, f in enumerate(archivos):
                ext = f.name.split('.')[-1].lower() if '.' in f.name else ''
                nuevo_nombre = f'SIN_{code_limpio}_{dt.now().timestamp()}_{i}.{ext}'
                DocumentoSinistro.objects.create(
                    code_sinistre=sinistro.code_sinistre,
                    nombre_archivo=nuevo_nombre,
                    tipo_documento=int(id_tipo) if id_tipo.isdigit() else None,
                    ruta_archivo=upload_dir + nuevo_nombre,
                    archivo=f,
                    sinistro=sinistro,
                )
                count += 1

        for doc_id in data.getlist('docs_reutilizados'):
            try:
                orig = DocumentoUsuario.objects.get(pk=int(doc_id))
            except (DocumentoUsuario.DoesNotExist, ValueError):
                continue
            DocumentoSinistro.objects.create(
                code_sinistre=sinistro.code_sinistre,
                nombre_archivo=orig.nombre_documento_usuario,
                tipo_documento=orig.tipo_documento_id,
                id_documento_usuario_origen=orig.id,
                sinistro=sinistro,
            )
            count += 1

        is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
        if is_ajax:
            return JsonResponse({'ok': True, 'count': count, 'id': sinistro.id})
        messages.success(request, f'{count} documento(s) añadido(s) al siniestro {sinistro.code_sinistre}')
        return redirect('sinistro_update', pk=sinistro.id)
    return redirect('sinistro_list')


def _resolver_documento_siniestro_legacy(nombre):
    if not nombre:
        return None
    for root in getattr(settings, 'DOCUMENTOS_LEGACY_ROOTS', []):
        candidato = os.path.join(str(root), nombre)
        if os.path.isfile(candidato):
            return candidato
    return None


@staff_member_required
def descargar_documento_siniestro(request, pk):
    doc = get_object_or_404(DocumentoSinistro, pk=pk)

    if doc.archivo:
        ruta = doc.archivo.path
        if os.path.isfile(ruta):
            response = FileResponse(open(ruta, 'rb'), as_attachment=True,
                                    filename=os.path.basename(ruta))
            return response

    if doc.ruta_archivo:
        ruta = os.path.join(str(settings.MEDIA_ROOT), doc.ruta_archivo)
        if os.path.isfile(ruta):
            response = FileResponse(open(ruta, 'rb'), as_attachment=True,
                                    filename=os.path.basename(ruta))
            return response

    ruta = _resolver_documento_siniestro_legacy(doc.nombre_archivo)
    if ruta:
        response = FileResponse(open(ruta, 'rb'), as_attachment=True,
                                filename=os.path.basename(ruta))
        return response

    return HttpResponseNotFound('Fichero no disponible.')



class SinistroUpdateView(View):
    template_name = 'siniestros/form.html'

    def get_object(self):
        return get_object_or_404(Sinistro, pk=self.kwargs['pk'])

    def get(self, request, pk):
        sinistro = self.get_object()
        poliza = sinistro.poliza
        val = {
            'code_sinistre': sinistro.code_sinistre or '',
            'ref_sinistre': sinistro.ref_sinistre or '',
            'cod_sig_correcto': '1' if sinistro.cod_sig_correcto else '0',
            'completo': '1' if sinistro.completo else '',
            'date_operation': sinistro.date_operation.strftime('%Y-%m-%d') if sinistro.date_operation else '',
            'date_sinistre': sinistro.date_sinistre.strftime('%Y-%m-%d') if sinistro.date_sinistre else '',
            'date_reception': sinistro.date_reception.strftime('%Y-%m-%d') if sinistro.date_reception else '',
            'police': sinistro.police or '',
            'assure': sinistro.assure or '',
            'categorie_assurance': sinistro.categorie_assurance or '',
            'reference_assurance': sinistro.reference_assurance or '',
            'benificiaire': sinistro.benificiaire or '',
            'reference_assistance': sinistro.reference_assistance or '',
            'nom_formule': sinistro.nom_formule or '',
            'nom_expert': sinistro.nom_expert or '',
            'reference_reglement': sinistro.reference_reglement or '',
            'document_de_base': sinistro.document_de_base or '',
            'montant': str(sinistro.montant) if sinistro.montant else '',
            'reference_payement': sinistro.reference_payement or '',
            'date_payement': sinistro.date_payement.strftime('%Y-%m-%d') if sinistro.date_payement else '',
            'id_recu': str(poliza.recu) if poliza else '',
        }
        return render(request, self.template_name, {
            'val': val,
            'sugerido': sinistro.code_sinistre,
            'formules': Formule.objects.all(),
            'expertos': Experto.objects.all(),
            'referencias_reg': ReferenceReglement.objects.all(),
            'documentos_base': DocumentDeBase.objects.all(),
            'tipos_doc': TipoDocumento.objects.all(),
            'hoy': datetime.now().strftime('%Y-%m-%d'),
            'editar': True,
            'poliza_sel': poliza,
        })

    def post(self, request, pk):
        sinistro = self.get_object()
        data = request.POST
        files = request.FILES

        poliza = None
        id_recu = data.get('id_recu')
        if id_recu and id_recu.isdigit():
            poliza = Poliza.objects.filter(recu=int(id_recu)).first()
        elif sinistro.poliza:
            poliza = sinistro.poliza

        sinistro.code_sinistre = data.get('code_sinistre', sinistro.code_sinistre)
        sinistro.ref_sinistre = data.get('ref_sinistre', sinistro.ref_sinistre)
        sinistro.date_operation = data.get('date_operation') or sinistro.date_operation
        sinistro.date_sinistre = data.get('date_sinistre') or sinistro.date_sinistre
        sinistro.date_reception = data.get('date_reception') or sinistro.date_reception
        sinistro.police = data.get('police', sinistro.police)
        sinistro.nom_formule = data.get('nom_formule', sinistro.nom_formule)
        sinistro.assure = data.get('assure', sinistro.assure)
        sinistro.categorie_assurance = data.get('categorie_assurance', sinistro.categorie_assurance)
        sinistro.reference_assurance = data.get('reference_assurance', sinistro.reference_assurance)
        sinistro.benificiaire = data.get('benificiaire', sinistro.benificiaire)
        sinistro.reference_assistance = data.get('reference_assistance', sinistro.reference_assistance)
        sinistro.nom_expert = data.get('nom_expert', sinistro.nom_expert)
        sinistro.reference_reglement = data.get('reference_reglement', sinistro.reference_reglement)
        sinistro.document_de_base = data.get('document_de_base', sinistro.document_de_base)
        sinistro.reference_payement = data.get('reference_payement', sinistro.reference_payement)
        sinistro.montant = float(data.get('montant', 0)) if data.get('montant') else None
        sinistro.date_payement = data.get('date_payement') or sinistro.date_payement
        sinistro.cod_sig_correcto = data.get('cod_sig_correcto') == '1'
        sinistro.completo = data.get('completo') == '1'
        sinistro.poliza = poliza
        sinistro.save()

        code_limpio = sinistro.code_sinistre.replace('/', '-') if sinistro.code_sinistre else f'SIN-{sinistro.id}'
        upload_dir = 'uploads_sinistres/'

        tipos_doc_up = data.getlist('tipo_doc_siniestro[]')
        for idx, id_tipo in enumerate(tipos_doc_up):
            if not id_tipo:
                continue
            file_key = f'archivos_siniestro_{idx}'
            archivos = request.FILES.getlist(file_key)
            for i, f in enumerate(archivos):
                ext = f.name.split('.')[-1].lower() if '.' in f.name else ''
                nuevo_nombre = f'SIN_{code_limpio}_{datetime.now().timestamp()}_{i}.{ext}'
                DocumentoSinistro.objects.create(
                    code_sinistre=sinistro.code_sinistre,
                    nombre_archivo=nuevo_nombre,
                    tipo_documento=int(id_tipo) if id_tipo.isdigit() else None,
                    ruta_archivo=upload_dir + nuevo_nombre,
                    archivo=f,
                    sinistro=sinistro,
                )

        for doc_id in data.getlist('docs_reutilizados'):
            try:
                orig = DocumentoUsuario.objects.get(pk=int(doc_id))
            except (DocumentoUsuario.DoesNotExist, ValueError):
                continue
            DocumentoSinistro.objects.create(
                code_sinistre=sinistro.code_sinistre,
                nombre_archivo=orig.nombre_documento_usuario,
                tipo_documento=orig.tipo_documento_id,
                id_documento_usuario_origen=orig.id,
                sinistro=sinistro,
            )

        messages.success(request, f'Siniestro {sinistro.code_sinistre} actualizado correctamente.')
        return redirect('sinistro_list')


def sinistro_delete(request, pk):
    sinistro = get_object_or_404(Sinistro, pk=pk)
    if request.method == 'POST':
        code = sinistro.code_sinistre
        sinistro.delete()
        messages.success(request, f'Siniestro {code} eliminado correctamente.')
    return redirect('sinistro_list')


def actualizar_completo_siniestro(request, pk):
    if request.method != 'POST':
        return JsonResponse({'status': 'error', 'mensaje': 'Método no permitido'}, status=405)
    sinistro = get_object_or_404(Sinistro, pk=pk)
    valor = request.POST.get('completo')
    completo = valor in ('1', 'true', 'True', 'on')
    sinistro.completo = completo
    sinistro.save(update_fields=['completo'])
    return JsonResponse({'status': 'ok', 'actualizado': True, 'completo': completo})


def sinistro_bulk_delete(request):
    if request.method == 'POST':
        pks = request.POST.getlist('pk')
        qs = Sinistro.objects.filter(pk__in=pks)
        count = qs.count()
        qs.delete()
        if count == 1:
            messages.success(request, 'Siniestro eliminado correctamente.')
        elif count > 1:
            messages.success(request, f'{count} siniestros eliminados correctamente.')
        else:
            messages.error(request, 'No se eliminó ningún siniestro.')
    return redirect('sinistro_list')


def sinistro_pdf(request, pk):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    import io

    sinistro = get_object_or_404(Sinistro, pk=pk)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph(f"<b>Siniestro: {sinistro.code_sinistre or '-'}</b>", styles['Title']))
    elements.append(Spacer(1, 6 * mm))

    fields = [
        ('Código Siniestro', sinistro.code_sinistre or '-'),
        ('Ref. Siniestro', sinistro.ref_sinistre or '-'),
        ('Fecha Operación', sinistro.date_operation.strftime('%d/%m/%Y') if sinistro.date_operation else '-'),
        ('Fecha Siniestro', sinistro.date_sinistre.strftime('%d/%m/%Y') if sinistro.date_sinistre else '-'),
        ('Fecha Recepción', sinistro.date_reception.strftime('%d/%m/%Y') if sinistro.date_reception else '-'),
        ('Póliza', sinistro.police or '-'),
        ('Asegurado', sinistro.assure or '-'),
        ('Categoría', sinistro.categorie_assurance or '-'),
        ('Ref. Assurance', sinistro.reference_assurance or '-'),
        ('Beneficiario', sinistro.benificiaire or '-'),
        ('Ref. Asistencia', sinistro.reference_assistance or '-'),
        ('Fórmula', sinistro.nom_formule or '-'),
        ('Experto', sinistro.nom_expert or '-'),
        ('Ref. Reglamento', sinistro.reference_reglement or '-'),
        ('Documento Base', sinistro.document_de_base or '-'),
        ('Ref. Pago', sinistro.reference_payement or '-'),
        ('Monto', f"{sinistro.montant:.2f} MAD" if sinistro.montant is not None else '-'),
        ('Fecha Pago', sinistro.date_payement.strftime('%d/%m/%Y') if sinistro.date_payement else '-'),
        ('Código Correcto', 'Sí' if sinistro.cod_sig_correcto else 'No'),
    ]

    data = [[Paragraph(f'<b>{k}</b>', styles['Normal']), Paragraph(v, styles['Normal'])] for k, v in fields]
    t = Table(data, colWidths=[120, 300])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.Color(0.9, 0.9, 0.9)),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(t)

    if sinistro.documentos.exists():
        elements.append(Spacer(1, 6 * mm))
        elements.append(Paragraph('<b>Documentos</b>', styles['Heading3']))
        doc_data = [['Archivo', 'Tipo']]
        for d in sinistro.documentos.all():
            doc_data.append([d.nombre_archivo or '-', str(d.tipo_documento) or '-'])
        t2 = Table(doc_data, colWidths=[250, 170])
        t2.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.Color(0.8, 0.8, 0.8)),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ]))
        elements.append(t2)

    doc.build(elements)
    buffer.seek(0)

    nombre_archivo = f"sinistro_{sinistro.code_sinistre or sinistro.id}.pdf".replace('/', '-')
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{nombre_archivo}"'
    return response


def sinistro_list_pdf(request):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    import io

    siniestros = Sinistro.objects.all()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph('<b>Lista de Siniestros</b>', styles['Title']))
    elements.append(Spacer(1, 6 * mm))
    elements.append(Paragraph(f'Total: {siniestros.count()} siniestros', styles['Normal']))
    elements.append(Spacer(1, 4 * mm))

    header = ['ID', 'Código', 'Fecha', 'Asegurado', 'Póliza', 'Monto', 'Pago']
    data = [header]
    for s in siniestros:
        data.append([
            str(s.id),
            s.code_sinistre or '-',
            s.date_sinistre.strftime('%d/%m/%Y') if s.date_sinistre else '-',
            s.assure or '-',
            s.police or '-',
            f"{s.montant:.2f} MAD" if s.montant is not None else '-',
            s.date_payement.strftime('%d/%m/%Y') if s.date_payement else '-',
        ])

    t = Table(data, colWidths=[40, 70, 70, 140, 100, 70, 70])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.Color(0.8, 0.8, 0.8)),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(t)

    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="lista_siniestros.pdf"'
    return response


def sinistro_list_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Siniestros'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='343a40', end_color='343a40', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    headers = ['ID', 'Código Siniestro', 'Ref. Siniestro', 'Fecha Siniestro', 'Asegurado',
               'Póliza', 'Categoría', 'Monto', 'Fecha Pago', 'Código Correcto']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    siniestros = Sinistro.objects.all()
    for i, s in enumerate(siniestros, 2):
        ws.cell(row=i, column=1, value=s.id).border = thin_border
        ws.cell(row=i, column=2, value=s.code_sinistre or '').border = thin_border
        ws.cell(row=i, column=3, value=s.ref_sinistre or '').border = thin_border
        ws.cell(row=i, column=4, value=s.date_sinistre.strftime('%d/%m/%Y') if s.date_sinistre else '').border = thin_border
        ws.cell(row=i, column=5, value=s.assure or '').border = thin_border
        ws.cell(row=i, column=6, value=s.police or '').border = thin_border
        ws.cell(row=i, column=7, value=s.categorie_assurance or '').border = thin_border
        ws.cell(row=i, column=8, value=float(s.montant) if s.montant is not None else '').border = thin_border
        ws.cell(row=i, column=9, value=s.date_payement.strftime('%d/%m/%Y') if s.date_payement else '').border = thin_border
        ws.cell(row=i, column=10, value='Sí' if s.cod_sig_correcto else 'No').border = thin_border

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 20

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="lista_siniestros.xlsx"'
    wb.save(response)
    return response
