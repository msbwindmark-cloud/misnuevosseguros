from django.urls import reverse_lazy
from django.views.generic import ListView, CreateView, UpdateView, DeleteView
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse, JsonResponse, HttpResponseNotFound, FileResponse
from django.contrib.admin.views.decorators import staff_member_required
from django.utils.decorators import method_decorator
from django.contrib import messages
from django.views import View
from datetime import date
from .models import EnvioCourrier, Destinataire, Expediteur
from .forms import DestinataireForm, ExpediteurForm, EnvioCourrierForm
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.units import cm
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl import Workbook

# --- Destinataire CRUD ---
@method_decorator(staff_member_required, name='dispatch')
class DestinataireListView(ListView):
    model = Destinataire
    template_name = 'courrier/destinataire_list.html'
    context_object_name = 'destinataires'

@method_decorator(staff_member_required, name='dispatch')
class DestinataireCreateView(CreateView):
    model = Destinataire
    form_class = DestinataireForm
    template_name = 'courrier/destinataire_form.html'
    success_url = reverse_lazy('destinataire_list')

@method_decorator(staff_member_required, name='dispatch')
class DestinataireUpdateView(UpdateView):
    model = Destinataire
    form_class = DestinataireForm
    template_name = 'courrier/destinataire_form.html'
    success_url = reverse_lazy('destinataire_list')

@method_decorator(staff_member_required, name='dispatch')
class DestinataireDeleteView(DeleteView):
    model = Destinataire
    template_name = 'courrier/destinataire_confirm_delete.html'
    success_url = reverse_lazy('destinataire_list')

# --- Expediteur CRUD ---
@method_decorator(staff_member_required, name='dispatch')
class ExpediteurListView(ListView):
    model = Expediteur
    template_name = 'courrier/expediteur_list.html'
    context_object_name = 'expediteurs'

@method_decorator(staff_member_required, name='dispatch')
class ExpediteurCreateView(CreateView):
    model = Expediteur
    form_class = ExpediteurForm
    template_name = 'courrier/expediteur_form.html'
    success_url = reverse_lazy('expediteur_list')

@method_decorator(staff_member_required, name='dispatch')
class ExpediteurUpdateView(UpdateView):
    model = Expediteur
    form_class = ExpediteurForm
    template_name = 'courrier/expediteur_form.html'
    success_url = reverse_lazy('expediteur_list')

@method_decorator(staff_member_required, name='dispatch')
class ExpediteurDeleteView(DeleteView):
    model = Expediteur
    template_name = 'courrier/expediteur_confirm_delete.html'
    success_url = reverse_lazy('expediteur_list')

# --- EnvioCourrier CRUD ---
class EnvioCourrierListView(ListView):
    model = EnvioCourrier
    template_name = 'courrier/lista.html'
    context_object_name = 'envios'

@method_decorator(staff_member_required, name='dispatch')
class EnvioCourrierCreateView(View):
    template_name = 'courrier/envio_form.html'
    success_url = reverse_lazy('courrier_list')

    def get(self, request):
        form = EnvioCourrierForm()
        expediteurs = Expediteur.objects.all()
        destinataires = Destinataire.objects.all()
        return render(request, self.template_name, {
            'form': form, 'expediteurs': expediteurs, 'destinataires': destinataires,
            'today': date.today(),
        })

    def post(self, request):
        data = request.POST
        files = request.FILES
        expediteur_id = data.get('expediteur')
        destinataire_id = data.get('destinataire')
        try:
            exp = Expediteur.objects.get(pk=int(expediteur_id)) if expediteur_id else None
        except (Expediteur.DoesNotExist, ValueError, TypeError):
            exp = None
        try:
            des = Destinataire.objects.get(pk=int(destinataire_id)) if destinataire_id else None
        except (Destinataire.DoesNotExist, ValueError, TypeError):
            des = None

        envio = EnvioCourrier.objects.create(
            date_envoie_courrier=data.get('date_envoie_courrier') or None,
            expediteur_envoie_courrier=exp.nom_expediteur if exp else data.get('expediteur_envoie_courrier', ''),
            addresse_expediteur_envoie_courrier=exp.adresse_expediteur if exp else data.get('addresse_expediteur_envoie_courrier', ''),
            destinataire_envoie_courrier=des.nom_destinataire if des else data.get('destinataire_envoie_courrier', ''),
            addresse_envoie_courrier=des.adresse_destinataire if des else data.get('addresse_envoie_courrier', ''),
        )

        file = files.get('fileupload')
        if file:
            ext = file.name.split('.')[-1].lower() if '.' in file.name else ''
            import os
            from django.conf import settings
            upload_dir = settings.MEDIA_ROOT / 'uploads_courrier'
            upload_dir.mkdir(parents=True, exist_ok=True)
            filepath = upload_dir / file.name
            with open(str(filepath), 'wb+') as dest:
                for chunk in file.chunks():
                    dest.write(chunk)
            envio.tipo_documento_envoie_courrier = ext
            envio.nombre_documento_envoie_courrier = file.name
            envio.tamanio_documento_envoie_courrier = file.size
            envio.save()

        messages.success(request, 'Envoie courrier enregistré correctement')
        return redirect(self.success_url)

@method_decorator(staff_member_required, name='dispatch')
class EnvioCourrierUpdateView(UpdateView):
    model = EnvioCourrier
    form_class = EnvioCourrierForm
    template_name = 'courrier/envio_form.html'
    success_url = reverse_lazy('courrier_list')

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['expediteurs'] = Expediteur.objects.all()
        ctx['destinataires'] = Destinataire.objects.all()
        ctx['today'] = date.today()  # <-- Añade esta línea aquí
        return ctx

    def form_valid(self, form):
        self.object = form.save(commit=False)
        exp_id = self.request.POST.get('expediteur')
        des_id = self.request.POST.get('destinataire')
        try:
            exp = Expediteur.objects.get(pk=int(exp_id)) if exp_id else None
        except (Expediteur.DoesNotExist, ValueError, TypeError):
            exp = None
        try:
            des = Destinataire.objects.get(pk=int(des_id)) if des_id else None
        except (Destinataire.DoesNotExist, ValueError, TypeError):
            des = None
        if exp:
            self.object.expediteur_envoie_courrier = exp.nom_expediteur
            self.object.addresse_expediteur_envoie_courrier = exp.adresse_expediteur
        if des:
            self.object.destinataire_envoie_courrier = des.nom_destinataire
            self.object.addresse_envoie_courrier = des.adresse_destinataire
        self.object.save()

        # 1. Si el usuario marcó la casilla de eliminar documento
        if self.request.POST.get('eliminar_documento'):
            import os
            from django.conf import settings
            old_fname = self.object.nombre_documento_envoie_courrier
            if old_fname:
                old_path = os.path.join(str(settings.MEDIA_ROOT), 'uploads_courrier', old_fname)
                if os.path.isfile(old_path):
                    os.remove(old_path)
            self.object.nombre_documento_envoie_courrier = ''
            self.object.tipo_documento_envoie_courrier = ''
            self.object.tamanio_documento_envoie_courrier = None
            self.object.save()

        # 2. Procesar nuevo archivo adjunto al editar (si se subió uno nuevo)
        file = self.request.FILES.get('fileupload')
        if file:
            ext = file.name.split('.')[-1].lower() if '.' in file.name else ''
            import os
            from django.conf import settings
            upload_dir = os.path.join(str(settings.MEDIA_ROOT), 'uploads_courrier')
            os.makedirs(upload_dir, exist_ok=True)
            filepath = os.path.join(upload_dir, file.name)
            with open(filepath, 'wb+') as dest:
                for chunk in file.chunks():
                    dest.write(chunk)
            self.object.tipo_documento_envoie_courrier = ext
            self.object.nombre_documento_envoie_courrier = file.name
            self.object.tamanio_documento_envoie_courrier = file.size
            self.object.save()

        return redirect(self.success_url)


@staff_member_required
def envio_courrier_download(request, pk):
    from django.conf import settings
    import os
    envio = get_object_or_404(EnvioCourrier, pk=pk)
    fname = getattr(envio, 'nombre_documento_envoie_courrier', '') or ''
    if not fname:
        return HttpResponseNotFound('No hay documento asociado con este courrier.')
    
    path = os.path.join(str(settings.MEDIA_ROOT), 'uploads_courrier', fname)
    if not os.path.isfile(path):
        return HttpResponseNotFound('El fichero no existe en el servidor.')
    
    f = open(path, 'rb')
    return FileResponse(f, as_attachment=True, filename=os.path.basename(fname))

@method_decorator(staff_member_required, name='dispatch')
class EnvioCourrierDeleteView(DeleteView):
    model = EnvioCourrier
    template_name = 'courrier/envio_confirm_delete.html'
    success_url = reverse_lazy('courrier_list')

@staff_member_required
def envio_courrier_last_pdf(request):
    envio = EnvioCourrier.objects.first()
    if not envio:
        return HttpResponse('No hay envíos de courrier', status=404)
    return envio_courrier_pdf(request, envio.pk)


@staff_member_required
def envio_courrier_pdf(request, pk):
    from reportlab.pdfgen import canvas as canvasmod
    from reportlab.lib.units import mm
    from .models import EnvioCourrier
    from datetime import datetime

    envio = EnvioCourrier.objects.get(pk=pk)
    buffer = BytesIO()
    w, h = A4
    c = canvasmod.Canvas(buffer, pagesize=A4)
    c.setTitle('Envoie courrier')
    c.setAuthor('SIS')

    x = 15 * mm
    y = h - 30 * mm
    c.setFont('Helvetica', 12)

    c.drawString(x, y, f'Recu No : {envio.id}')
    y -= 10 * mm
    c.drawString(x, y, f'Date Envoie courrier      :      {envio.date_envoie_courrier.strftime("%d/%m/%Y") if envio.date_envoie_courrier else ""}')
    y -= 8 * mm
    c.drawString(x, y, f'Expediteur envoie courrier      :      {envio.expediteur_envoie_courrier}')
    y -= 14 * mm
    c.drawString(x, y, f'Destinataire envoie courrier         :      {envio.destinataire_envoie_courrier}')
    y -= 7 * mm
    c.drawString(x, y, f'Addresse envoie courrier        :     {envio.addresse_envoie_courrier}')
    y -= 7 * mm
    fichero = envio.nombre_documento_envoie_courrier or ''
    c.drawString(x, y, f'Fichiers envoie courrier        :     {fichero}')
    y -= 14 * mm
    c.setFont('Helvetica', 10)
    c.drawString(x, y, f'Document: Justificatif courrier N {envio.id}')
    y -= 5 * mm
    c.drawString(x, y, f'Imprime le: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}')

    c.save()
    buffer.seek(0)
    nombre = f"justificatif_courrier_{envio.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return HttpResponse(buffer, content_type='application/pdf', headers={
        'Content-Disposition': f'attachment; filename="{nombre}"'
    })

@staff_member_required
def envio_courrier_list_pdf(request):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []
    elements.append(Paragraph('Lista de Envios de Courrier', styles['Title']))
    elements.append(Spacer(1, 0.5*cm))
    data = [['ID', 'Fecha', 'Expediteur', 'Destinataire', 'Adresse']]
    for e in EnvioCourrier.objects.all().order_by('-id'):
        data.append([
            str(e.id),
            e.date_envoie_courrier.strftime('%d/%m/%Y') if e.date_envoie_courrier else '',
            e.expediteur_envoie_courrier,
            e.destinataire_envoie_courrier,
            e.addresse_envoie_courrier,
        ])
    table = Table(data, colWidths=[2*cm, 2.5*cm, 4*cm, 4*cm, 5*cm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#4e73df')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('GRID', (0,0), (-1,-1), 1, colors.black),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    return HttpResponse(buffer, content_type='application/pdf', headers={
        'Content-Disposition': 'attachment; filename="lista_courrier.pdf"'
    })


@staff_member_required
def export_envios_courrier_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Envíos Courrier'
    headers = ['ID', 'Fecha', 'Remitente', 'Destinatario', 'Dirección', 'Tipo Doc', 'Nombre Doc', 'Tamaño', 'Creado']
    header_fill = PatternFill(start_color='198754', end_color='198754', fill_type='solid')
    header_font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal='center')
    for row, e in enumerate(EnvioCourrier.objects.all().order_by('-id'), 2):
        ws.cell(row=row, column=1, value=e.id)
        ws.cell(row=row, column=2, value=e.date_envoie_courrier.strftime('%d/%m/%Y') if e.date_envoie_courrier else '')
        ws.cell(row=row, column=3, value=e.expediteur_envoie_courrier)
        ws.cell(row=row, column=4, value=e.destinataire_envoie_courrier)
        ws.cell(row=row, column=5, value=e.addresse_envoie_courrier)
        ws.cell(row=row, column=6, value=e.tipo_documento_envoie_courrier or '')
        ws.cell(row=row, column=7, value=e.nombre_documento_envoie_courrier or '')
        ws.cell(row=row, column=8, value=e.tamanio_documento_envoie_courrier or '')
        ws.cell(row=row, column=9, value=e.creado_el.strftime('%d/%m/%Y %H:%M') if e.creado_el else '')
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 35
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 20
    ws.column_dimensions['H'].width = 12
    ws.column_dimensions['I'].width = 18
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="envios_courrier.xlsx"'
    wb.save(response)
    return response


@staff_member_required
def get_expediteur_json(request):
    pk = request.GET.get('pk', '')
    if not pk or not pk.isdigit():
        return JsonResponse({'ok': False})
    try:
        e = Expediteur.objects.get(pk=int(pk))
        return JsonResponse({'ok': True, 'nom': e.nom_expediteur, 'addresse': e.adresse_expediteur})
    except Expediteur.DoesNotExist:
        return JsonResponse({'ok': False})


@staff_member_required
def get_destinataire_json(request):
    pk = request.GET.get('pk', '')
    if not pk or not pk.isdigit():
        return JsonResponse({'ok': False})
    try:
        d = Destinataire.objects.get(pk=int(pk))
        return JsonResponse({'ok': True, 'nom': d.nom_destinataire, 'addresse': d.adresse_destinataire})
    except Destinataire.DoesNotExist:
        return JsonResponse({'ok': False})


# --- AJAX Expediteur (generer_expediteur.php) ---

@method_decorator(staff_member_required, name='dispatch')
class ExpediteurAjaxView(View):
    def post(self, request):
        action = request.POST.get('action', '')
        if action == 'insert':
            Expediteur.objects.create(
                nom_expediteur=request.POST.get('nom_expediteur', ''),
                adresse_expediteur=request.POST.get('adresse_expediteur', ''),
            )
            return JsonResponse({'ok': True})
        if request.POST.get('edit_id'):
            e = Expediteur.objects.get(pk=request.POST.get('edit_id'))
            return JsonResponse({'id': e.pk, 'nom_expediteur': e.nom_expediteur, 'adresse_expediteur': e.adresse_expediteur})
        if action == 'update':
            e = Expediteur.objects.get(pk=request.POST.get('id_expediteur'))
            e.nom_expediteur = request.POST.get('nom_expediteur', e.nom_expediteur)
            e.adresse_expediteur = request.POST.get('adresse_expediteur', e.adresse_expediteur)
            e.save()
            return JsonResponse({'ok': True})
        if request.POST.get('del_id'):
            Expediteur.objects.get(pk=request.POST.get('del_id')).delete()
            return JsonResponse({'ok': True})
        return JsonResponse({'ok': False, 'error': 'Acción no reconocida'})


# --- AJAX Destinataire (generer_destinataire.php) ---

@method_decorator(staff_member_required, name='dispatch')
class DestinataireAjaxView(View):
    def post(self, request):
        action = request.POST.get('action', '')
        if action == 'insert':
            Destinataire.objects.create(
                nom_destinataire=request.POST.get('nom_destinataire', ''),
                adresse_destinataire=request.POST.get('adresse_destinataire', ''),
            )
            return JsonResponse({'ok': True})
        if request.POST.get('edit_id'):
            d = Destinataire.objects.get(pk=request.POST.get('edit_id'))
            return JsonResponse({'id': d.pk, 'nom_destinataire': d.nom_destinataire, 'adresse_destinataire': d.adresse_destinataire})
        if action == 'update':
            d = Destinataire.objects.get(pk=request.POST.get('id_destinataire'))
            d.nom_destinataire = request.POST.get('nom_destinataire', d.nom_destinataire)
            d.adresse_destinataire = request.POST.get('adresse_destinataire', d.adresse_destinataire)
            d.save()
            return JsonResponse({'ok': True})
        if request.POST.get('del_id'):
            Destinataire.objects.get(pk=request.POST.get('del_id')).delete()
            return JsonResponse({'ok': True})
        return JsonResponse({'ok': False, 'error': 'Acción no reconocida'})
